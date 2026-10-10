import openpyxl, subprocess, json, re, os, datetime

DB_NAME = "internship-registration-system-db"

def run_cmd(cmd):
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if res.returncode != 0:
        raise Exception(f"Command failed: {' '.join(cmd)}\nStderr: {res.stderr}\nStdout: {res.stdout}")
    return res.stdout

def d1_query(sql):
    cmd = ["npx", "wrangler", "d1", "execute", DB_NAME, "--remote", f"--command={sql}", "--json"]
    out = run_cmd(cmd)
    try:
        idx = out.find("[")
        if idx != -1:
            data = json.loads(out[idx:])
            return data[0]["results"]
        raise Exception("No JSON array found in output")
    except Exception as e:
        raise Exception(f"Failed to parse D1 query output for {sql}: {e}\nRaw: {out}")

def d1_execute_file(filepath):
    cmd = ["npx", "wrangler", "d1", "execute", DB_NAME, "--remote", f"--file={filepath}"]
    out = run_cmd(cmd)
    print(f"Executed {filepath}: {out.strip()[:200]}...")
    return out

def sql_str(val):
    if val is None:
        return "NULL"
    s = str(val).strip()
    if s == "" or s.lower() == "none":
        return "NULL"
    escaped = s.replace("'", "''")
    return f"'{escaped}'"

def sql_num(val):
    if val is None:
        return "NULL"
    try:
        return str(float(val)) if "." in str(val) else str(int(val))
    except:
        return "NULL"

def parse_vietnamese_date(dt_str):
    if not dt_str:
        return "CURRENT_TIMESTAMP"
    s = str(dt_str).strip()
    # Format e.g. "20:54:03 13/7/2026" or "10:30:39 26/5/2026"
    m = re.match(r"(\d{1,2}):(\d{1,2}):(\d{1,2})\s+(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        hh, mm, ss, dd, mo, yyyy = m.groups()
        return f"'{yyyy}-{int(mo):02d}-{int(dd):02d} {int(hh):02d}:{int(mm):02d}:{int(ss):02d}'"
    return "CURRENT_TIMESTAMP"

def main():
    print("=== STEP 1: LOAD SOURCE FILES ===")
    comp_file = "/Users/tuyenkv/Downloads/danh_sach_cong_ty.xlsx"
    reg_file = "/Users/tuyenkv/Downloads/danh_sach_dang_ky_thuc_tap.xlsx"
    cur_file = "/Users/tuyenkv/Downloads/danh_sach_hien_tai (1).xlsx"
    grad_file = "/Users/tuyenkv/Downloads/bang_diem_thuc_tap.xlsx"

    wb_comp = openpyxl.load_workbook(comp_file)
    comp_rows = list(wb_comp.active.iter_rows(values_only=True))
    h_c = comp_rows[0]

    wb_reg = openpyxl.load_workbook(reg_file)
    reg_rows = list(wb_reg.active.iter_rows(values_only=True))
    h_reg = reg_rows[0]

    wb_cur = openpyxl.load_workbook(cur_file)
    cur_rows = list(wb_cur.active.iter_rows(values_only=True))
    h_cur = cur_rows[0]

    wb_grad = openpyxl.load_workbook(grad_file)
    grad_rows = list(wb_grad.active.iter_rows(values_only=True))
    h_g = grad_rows[0]

    print(f"Loaded: {len(comp_rows)-1} companies, {len(reg_rows)-1} students, {len(cur_rows)-1} registrations, {len(grad_rows)-1} grades.")

    os.makedirs("scratch/d1_migration", exist_ok=True)

    # ----------------------------------------------------
    # 1. LECTURERS
    # ----------------------------------------------------
    print("\n=== STEP 2: ENSURE ALL LECTURERS ===")
    missing_lecturers = [
        'PGS.TS. Hà Quang Thuỵ', 'PGS.TS. Nguyễn Hoài Sơn', 'PGS.TS. Nguyễn Việt Hà',
        'PGS.TS. Phan Hải Đăng', 'PGS.TS. Trương Ninh Thuận', 'TS. Bùi Quang Hưng',
        'TS. Lê Trung Thành', 'ThS. Lương Việt Nguyên', 'ThS. Nguyễn Thị Thuỳ Linh',
        'ThS. Đỗ Hoàng Kiên'
    ]
    with open("scratch/d1_migration/01_lecturers.sql", "w", encoding="utf-8") as f:
        for l in missing_lecturers:
            f.write(f"INSERT OR IGNORE INTO lecturers (name) VALUES ({sql_str(l)});\n")
    d1_execute_file("scratch/d1_migration/01_lecturers.sql")

    # Fetch all lecturers
    d1_lecturers = d1_query("SELECT id, name FROM lecturers;")
    def strip_title(name):
        return re.sub(r'^(GS|PGS|TS|ThS|CN|Admin|tuyenkv)[\.\s]+', '', name, flags=re.IGNORECASE).strip().lower()

    lecturer_map_exact = {l["name"].strip(): l["id"] for l in d1_lecturers}
    lecturer_map_clean = {strip_title(l["name"]): l["id"] for l in d1_lecturers}

    def get_lecturer_id(name):
        if not name: return None
        n = str(name).strip()
        if n in lecturer_map_exact: return lecturer_map_exact[n]
        c = strip_title(n)
        if c in lecturer_map_clean: return lecturer_map_clean[c]
        return None

    # ----------------------------------------------------
    # 2. COMPANIES
    # ----------------------------------------------------
    print("\n=== STEP 3: INSERT / UPDATE COMPANIES ===")
    # Query existing companies
    existing_companies = d1_query("SELECT id, name FROM companies;")
    existing_comp_map = {c["name"].strip(): c["id"] for c in existing_companies}

    # Prepare company statements
    comp_sql = []
    # Ensure "Công ty khác"
    if "Công ty khác" not in existing_comp_map:
        comp_sql.append("INSERT INTO companies (name, description, slots) VALUES ('Công ty khác', 'Sinh viên tự liên hệ nơi thực tập ngoài danh sách', 999);")

    # Ensure "Trường Đại học Công nghệ"
    if "Trường Đại học Công nghệ" not in existing_comp_map:
        comp_sql.append("INSERT INTO companies (name, description, slots) VALUES ('Trường Đại học Công nghệ', 'Thực tập tại các phòng thí nghiệm, trung tâm nghiên cứu thuộc Trường ĐH Công nghệ', 150);")

    for r in comp_rows[1:]:
        c_name = str(r[h_c.index('Tên doanh nghiệp')] or '').strip()
        if not c_name: continue
        loai = str(r[h_c.index('Loại')] or '').strip()
        slots = r[h_c.index('Chỉ tiêu')] or 5
        drive_link = r[h_c.index('Link Drive')]
        email = r[h_c.index('Email liên hệ')]
        contact = r[h_c.index('Người liên hệ')]
        phone = r[h_c.index('SĐT')]
        address = r[h_c.index('Địa chỉ')]

        if c_name in existing_comp_map:
            cid = existing_comp_map[c_name]
            comp_sql.append(
                f"UPDATE companies SET slots = {sql_num(slots)}, contact_email = {sql_str(email)}, "
                f"contact_name = {sql_str(contact)}, phone = {sql_str(phone)}, address = {sql_str(address)}, "
                f"applicants_drive_link = {sql_str(drive_link)} WHERE id = {cid};"
            )
        else:
            comp_sql.append(
                f"INSERT INTO companies (name, description, slots, contact_email, contact_name, phone, address, applicants_drive_link) "
                f"VALUES ({sql_str(c_name)}, {sql_str(loai + ' - ' + c_name)}, {sql_num(slots)}, {sql_str(email)}, "
                f"{sql_str(contact)}, {sql_str(phone)}, {sql_str(address)}, {sql_str(drive_link)});"
            )

    with open("scratch/d1_migration/02_companies.sql", "w", encoding="utf-8") as f:
        f.write("\n".join(comp_sql) + "\n")
    d1_execute_file("scratch/d1_migration/02_companies.sql")

    # Refresh companies map
    updated_companies = d1_query("SELECT id, name FROM companies;")
    comp_map = {c["name"].strip(): c["id"] for c in updated_companies}
    print(f"Total companies in D1 now: {len(comp_map)}")

    # ----------------------------------------------------
    # 3. USERS (STUDENTS)
    # ----------------------------------------------------
    print("\n=== STEP 4: UPSERT ALL STUDENTS ===")
    student_dict = {}
    for r in reg_rows[1:]:
        sid = str(r[h_reg.index('MSSV')]).strip()
        if not sid: continue
        student_dict[sid] = {
            'student_id': sid,
            'name': str(r[h_reg.index('Họ và tên')] or '').strip(),
            'dob': str(r[h_reg.index('Ngày sinh')] or '').strip(),
            'class_name': str(r[h_reg.index('Lớp khoá học')] or '').strip(),
            'course_code': str(r[h_reg.index('Mã môn học')] or '').strip(),
            'phone': str(r[h_reg.index('SĐT')] or '').strip() if r[h_reg.index('SĐT')] else '',
            'personal_email': str(r[h_reg.index('Email cá nhân')] or '').strip() if r[h_reg.index('Email cá nhân')] else '',
            'email': str(r[h_reg.index('Email VNU')] or '').strip() or f'{sid}@vnu.edu.vn',
        }

    # Add any extra students from danh_sach_hien_tai
    for r in cur_rows[1:]:
        sid = str(r[h_cur.index('Mã SV')]).strip()
        if not sid: continue
        if sid not in student_dict:
            student_dict[sid] = {
                'student_id': sid,
                'name': str(r[h_cur.index('Họ và tên')] or '').strip(),
                'dob': str(r[h_cur.index('Ngày sinh')] or '').strip(),
                'class_name': str(r[h_cur.index('Lớp KH')] or '').strip(),
                'course_code': str(r[h_cur.index('Mã môn học')] or '').strip(),
                'phone': str(r[h_cur.index('SĐT')] or '').strip() if r[h_cur.index('SĐT')] else '',
                'personal_email': str(r[h_cur.index('Email cá nhân')] or '').strip() if r[h_cur.index('Email cá nhân')] else '',
                'email': f'{sid}@vnu.edu.vn',
            }

    users_sql = []
    for sid, s in student_dict.items():
        users_sql.append(
            f"INSERT INTO users (email, name, role, student_id, dob, class_name, course_code, phone, personal_email) "
            f"VALUES ({sql_str(s['email'])}, {sql_str(s['name'])}, 'student', {sql_str(s['student_id'])}, "
            f"{sql_str(s['dob'])}, {sql_str(s['class_name'])}, {sql_str(s['course_code'])}, {sql_str(s['phone'])}, {sql_str(s['personal_email'])}) "
            f"ON CONFLICT(email) DO UPDATE SET name = excluded.name, student_id = excluded.student_id, dob = excluded.dob, "
            f"class_name = excluded.class_name, course_code = excluded.course_code, phone = excluded.phone, personal_email = excluded.personal_email "
            f"WHERE role = 'student';"
        )

    with open("scratch/d1_migration/03_users.sql", "w", encoding="utf-8") as f:
        f.write("\n".join(users_sql) + "\n")
    d1_execute_file("scratch/d1_migration/03_users.sql")

    # Fetch users map
    d1_users = d1_query("SELECT id, student_id, email FROM users;")
    user_by_sid = {u["student_id"].strip(): u["id"] for u in d1_users if u.get("student_id")}
    user_by_email = {u["email"].strip(): u["id"] for u in d1_users if u.get("email")}
    print(f"Total users in D1: {len(d1_users)}, student mappings: {len(user_by_sid)}")

    # ----------------------------------------------------
    # 4. REGISTRATIONS
    # ----------------------------------------------------
    print("\n=== STEP 5: INSERT REGISTRATIONS ===")
    # Clear old registrations & final_internships & grades & advisor_assignments to prevent any dangling state
    cleanup_sql = """
    DELETE FROM advisor_assignments;
    DELETE FROM grades;
    DELETE FROM final_internships;
    DELETE FROM registrations;
    """
    with open("scratch/d1_migration/04_cleanup.sql", "w", encoding="utf-8") as f:
        f.write(cleanup_sql)
    d1_execute_file("scratch/d1_migration/04_cleanup.sql")

    reg_sql_batches = []
    cur_batch = []

    # Map each registration
    student_nv_counter = {}
    for idx, r in enumerate(cur_rows[1:]):
        sid = str(r[h_cur.index('Mã SV')]).strip()
        uid = user_by_sid.get(sid) or user_by_email.get(f"{sid}@vnu.edu.vn")
        if not uid:
            print(f"Warning: user {sid} not found in users map!")
            continue

        student_nv_counter[sid] = student_nv_counter.get(sid, 0) + 1
        pref_order = student_nv_counter[sid]

        place_raw = str(r[h_cur.index('Nơi thực tập')] or '').strip()
        if place_raw.startswith("Công ty khác:"):
            cid = comp_map.get("Công ty khác")
            other_name = place_raw[len("Công ty khác:"):].strip()
            other_role = str(r[h_cur.index('Vị trí')] or '').strip()
            other_contact = str(r[h_cur.index('Liên hệ')] or '').strip()
            note = str(r[h_cur.index('Ghi chú')] or '').strip()
        elif place_raw == "Trường Đại học Công nghệ":
            cid = comp_map.get("Trường Đại học Công nghệ")
            other_name = ""
            other_role = ""
            other_contact = ""
            note = str(r[h_cur.index('Ghi chú')] or '').strip()
            # If note has GVHD, extract contact
            if note.startswith("GVHD:"):
                other_contact = note[5:].split(" - ")[0].strip()
        else:
            cid = comp_map.get(place_raw)
            if not cid:
                cid = comp_map.get("Công ty khác")
                other_name = place_raw
                other_role = str(r[h_cur.index('Vị trí')] or '').strip()
                other_contact = str(r[h_cur.index('Liên hệ')] or '').strip()
            else:
                other_name = ""
                other_role = str(r[h_cur.index('Vị trí')] or '').strip()
                other_contact = str(r[h_cur.index('Liên hệ')] or '').strip()
            note = str(r[h_cur.index('Ghi chú')] or '').strip()

        st_raw = str(r[h_cur.index('Trạng thái')] or '').strip()
        if st_raw == "Đã duyệt": status = "approved"
        elif st_raw == "Từ chối": status = "rejected"
        else: status = "pending"

        rev_comment = str(r[h_cur.index('Nhận xét duyệt')] or '').strip()
        sent_dn = parse_vietnamese_date(r[h_cur.index('Đã gửi DN')])
        created_at = parse_vietnamese_date(r[h_cur.index('Thời gian đăng ký')])

        sql = (
            f"INSERT INTO registrations (user_id, company_id, note, status, created_at, other_company_name, "
            f"other_company_role, other_company_contact, sent_to_company_at, review_comment, preference_order) "
            f"VALUES ({uid}, {cid}, {sql_str(note)}, {sql_str(status)}, {created_at}, {sql_str(other_name)}, "
            f"{sql_str(other_role)}, {sql_str(other_contact)}, {sent_dn if sent_dn != 'CURRENT_TIMESTAMP' else 'NULL'}, "
            f"{sql_str(rev_comment)}, {pref_order});"
        )
        cur_batch.append(sql)
        if len(cur_batch) >= 500:
            reg_sql_batches.append(cur_batch)
            cur_batch = []
    if cur_batch:
        reg_sql_batches.append(cur_batch)

    for b_idx, batch in enumerate(reg_sql_batches):
        fname = f"scratch/d1_migration/05_registrations_part{b_idx+1}.sql"
        with open(fname, "w", encoding="utf-8") as f:
            f.write("\n".join(batch) + "\n")
        d1_execute_file(fname)

    # ----------------------------------------------------
    # 5. FINAL INTERNSHIPS & GRADES & ADVISORS
    # ----------------------------------------------------
    print("\n=== STEP 6: INSERT FINAL INTERNSHIPS & GRADES & ADVISORS ===")
    # Query fresh registrations
    d1_regs = d1_query("SELECT id, user_id, company_id, other_company_name FROM registrations;")
    regs_by_user = {}
    for rg in d1_regs:
        regs_by_user.setdefault(rg["user_id"], []).append(rg)

    final_sql = []
    grades_sql = []
    advisor_sql = []

    for r in grad_rows[1:]:
        sid = str(r[h_g.index('Mã SV')]).strip()
        uid = user_by_sid.get(sid) or user_by_email.get(f"{sid}@vnu.edu.vn")
        if not uid:
            print(f"Warning: student {sid} not found for grade entry!")
            continue

        place_raw = str(r[h_g.index('Nơi thực tập')] or '').strip()
        gvhd_main = str(r[h_g.index('GVHD chính')] or '').strip()
        gvdhd = str(r[h_g.index('Đồng hướng dẫn')] or '').strip()
        score_prog = r[h_g.index('Điểm định kỳ')]
        score_rep = r[h_g.index('Điểm báo cáo')]
        score_comp = r[h_g.index('Điểm công ty/GVHD')]
        score_final = r[h_g.index('Điểm tổng kết')]
        st_grade = str(r[h_g.index('Trạng thái')] or '').strip()
        grader = str(r[h_g.index('Người nhập')] or '').strip()
        sub_time = parse_vietnamese_date(r[h_g.index('Nộp điểm lúc')])
        note_grade = str(r[h_g.index('Ghi chú')] or '').strip()

        # Determine company & type
        if place_raw == "Trường Đại học Công nghệ":
            cid = comp_map.get("Trường Đại học Công nghệ")
            itype = "school"
            other_name = ""
        elif place_raw in comp_map:
            cid = comp_map[place_raw]
            itype = "official"
            other_name = ""
        else:
            cid = comp_map.get("Công ty khác")
            itype = "external"
            other_name = place_raw.replace("Công ty khác:", "").strip()

        # Find matching registration
        matching_reg_id = None
        user_regs = regs_by_user.get(uid, [])
        for ur in user_regs:
            if ur["company_id"] == cid:
                if itype == "external":
                    if (ur.get("other_company_name") or "").lower() == other_name.lower():
                        matching_reg_id = ur["id"]
                        break
                else:
                    matching_reg_id = ur["id"]
                    break

        final_sql.append(
            f"INSERT INTO final_internships (user_id, registration_id, company_id, internship_type, status, "
            f"student_attested, school_lecturer, confirmed_at, note) "
            f"VALUES ({uid}, {sql_num(matching_reg_id)}, {cid}, {sql_str(itype)}, 'confirmed', 1, "
            f"{sql_str(gvhd_main)}, CURRENT_TIMESTAMP, {sql_str(note_grade)});"
        )

        # Grades
        lecturer_id = get_lecturer_id(gvhd_main) or get_lecturer_id(grader) or 1
        grade_status = "submitted" if st_grade == "Đã nộp" else ("draft" if score_final is not None else "draft")
        grades_sql.append(
            f"INSERT INTO grades (user_id, lecturer_id, progress_score, report_score, company_score, final_score, "
            f"status, comment, submitted_at) "
            f"VALUES ({uid}, {lecturer_id}, {sql_num(score_prog)}, {sql_num(score_rep)}, {sql_num(score_comp)}, "
            f"{sql_num(score_final)}, {sql_str(grade_status)}, {sql_str(note_grade)}, "
            f"{sub_time if sub_time != 'CURRENT_TIMESTAMP' else 'NULL'});"
        )

        # Advisor assignments
        if gvhd_main:
            for l_item in gvhd_main.split(";"):
                l_item = l_item.strip()
                lid = get_lecturer_id(l_item)
                if lid:
                    advisor_sql.append(f"INSERT OR IGNORE INTO advisor_assignments (user_id, lecturer_id, role) VALUES ({uid}, {lid}, 'primary');")

        if gvdhd:
            for l_item in gvdhd.split(";"):
                l_item = l_item.strip()
                lid = get_lecturer_id(l_item)
                if lid:
                    advisor_sql.append(f"INSERT OR IGNORE INTO advisor_assignments (user_id, lecturer_id, role) VALUES ({uid}, {lid}, 'co');")

    with open("scratch/d1_migration/06_final_internships.sql", "w", encoding="utf-8") as f:
        f.write("\n".join(final_sql) + "\n")
    d1_execute_file("scratch/d1_migration/06_final_internships.sql")

    with open("scratch/d1_migration/07_grades.sql", "w", encoding="utf-8") as f:
        f.write("\n".join(grades_sql) + "\n")
    d1_execute_file("scratch/d1_migration/07_grades.sql")

    with open("scratch/d1_migration/08_advisors.sql", "w", encoding="utf-8") as f:
        f.write("\n".join(advisor_sql) + "\n")
    d1_execute_file("scratch/d1_migration/08_advisors.sql")

    print("\n=== STEP 7: VERIFY COUNTS IN D1 ===")
    stats = d1_query("""
    SELECT 
      (SELECT count(*) FROM users) as users_count,
      (SELECT count(*) FROM companies) as companies_count,
      (SELECT count(*) FROM registrations) as registrations_count,
      (SELECT count(*) FROM final_internships) as final_internships_count,
      (SELECT count(*) FROM grades) as grades_count,
      (SELECT count(*) FROM advisor_assignments) as advisors_count;
    """)
    print("D1 Migration Results:", json.dumps(stats[0], indent=2))

if __name__ == "__main__":
    main()
